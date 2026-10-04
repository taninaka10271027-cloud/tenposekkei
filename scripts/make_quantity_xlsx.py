# -*- coding: utf-8 -*-
"""数量拾い出し → 数量表（Excel）。

入力: output/quantities.json（make_elec_plan.py が平面図の作図データから集計）
出力: output/数量表_電灯コンセント設備.xlsx

ケーブル長の拾い方:
  平面長  … 図上の配線長（器具中心間、直交経路）。ホームランは器具から分電盤までの直交距離
  立上下り … スイッチ 1.1m (CH2.4-FL1.3)、コンセント 2.1m (CH2.4-FL0.3)、分電盤 0.6m (盤上端 FL+1.8)
  余長    … (平面長+立上下り) × 10%
"""
import json
import math
import os
import sys
from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(__file__))
from circuits import PANELS, panel_rows, phase_sum  # noqa: E402

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
OUT = os.path.join(ROOT, "output")
Q = json.load(open(os.path.join(OUT, "quantities.json"), encoding="utf-8"))
DEV = Q["devices"]
SLACK = 0.10

thin = Side(style="thin", color="808080")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD = PatternFill("solid", fgColor="DDEBF7")
SUB = PatternFill("solid", fgColor="F2F2F2")
FONT = "游ゴシック"


def panel_of(circuit):
    if not circuit:
        return None
    return "L-1" if circuit.startswith("L1") else "L-2"


def style_table(ws, r0, r1, c1):
    for row in ws.iter_rows(min_row=r0, max_row=r1, max_col=c1):
        for c in row:
            c.border = BORDER
            c.font = Font(name=FONT, size=10, bold=c.row == r0)
            if c.row == r0:
                c.fill = HEAD
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                c.alignment = Alignment(vertical="center", wrap_text=isinstance(c.value, str) and len(c.value) > 30)


def widths(ws, ws_w):
    for i, w in enumerate(ws_w, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def title(ws, s, sub=None):
    ws["A1"] = s
    ws["A1"].font = Font(name=FONT, size=14, bold=True)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(name=FONT, size=9, color="555555")


# ---------------------------------------------------------------- 集計
FIX_NAME = {"DL": ("LEDダウンライト", "φ100 埋込 1,000lm 電球色 約8.5W"),
            "BL": ("LEDベースライト", "40形 逆富士 直付 6,900lm 昼白色 約45W")}
SW_NAME = {None: ("片切スイッチ", "15A 埋込 FL+1,300"), "3": ("3路スイッチ", "15A 埋込 FL+1,300"),
           "L": ("確認表示灯内蔵スイッチ", "15A 埋込 FL+1,300（便所）")}

count = defaultdict(Counter)      # (品名,仕様) -> Counter(panel)
room_count = defaultdict(Counter)  # room -> Counter(品名)
for d in DEV:
    p = panel_of(d.get("circuit"))
    if d["kind"] in FIX_NAME:
        k = FIX_NAME[d["kind"]]
    elif d["kind"] == "SW":
        k = SW_NAME[d.get("sub")]
    elif d["kind"] == "OUT":
        k = ("コンセント", d["spec"] + " FL+300")
    else:
        continue
    count[("器具・配線器具",) + k][p] += 1
    room_count[d["room"]][k[0] if d["kind"] != "OUT" else "コンセント（" + d["sym"] + "）"] += 1

# スイッチの設置箇所（連用）: 同じ壁に 0.3m 以内で並ぶものを 1 箇所とする
sws = [d for d in DEV if d["kind"] == "SW"]
groups = []
for d in sws:
    for g in groups:
        if any(abs(d["x"] - e["x"]) < 50 and abs(d["y"] - e["y"]) < 300 for e in g):
            g.append(d)
            break
    else:
        groups.append([d])
gang = Counter()
for g in groups:
    gang[(len(g), panel_of(g[0]["circuit"]))] += 1
for (n, p), c in gang.items():
    count[("ボックス・プレート", "スイッチボックス", "1個用（%d連）" % n)][p] += c
    count[("ボックス・プレート", "スイッチプレート", "%d口用" % n)][p] += c
outs = [d for d in DEV if d["kind"] == "OUT"]
for d in outs:
    p = panel_of(d["circuit"])
    if "防雨" in d["spec"]:
        count[("ボックス・プレート", "防雨形ボックス", "コンセント用")][p] += 1
    else:
        count[("ボックス・プレート", "スイッチボックス", "1個用（コンセント）")][p] += 1
        count[("ボックス・プレート", "コンセントプレート", "1口用" if d["sym"] in ("EET", "E20A", "E20A250V") else "2口用")][p] += 1

# ジョイントボックス: 配線の分岐点（器具の接続数 3 以上）を 1 箇所とする
deg = Counter()
for a, b in Q["edges"]:
    for e in (a, b):
        if e and e != "HOMERUN":
            deg[e] += 1
dev_by_id = {d["id"]: d for d in DEV}
for i, n in deg.items():
    d = dev_by_id[i]
    if n >= 3 and d["kind"] in ("DL", "BL", "OUT"):
        count[("ボックス・プレート", "ジョイントボックス", "VVF用 天井内")][panel_of(d["circuit"])] += 1

# ケーブル
cab = defaultdict(lambda: Counter())
cab_rows = []
for r in Q["cable_runs"]:
    p = panel_of(r["circuit"])
    plan, drop = r["plan"] / 1000, r["drop"] / 1000
    total = (plan + drop) * (1 + SLACK)
    cab_rows.append((r["circuit"], r["kind"], r["cable"], plan, drop, total))
    cables = [r["cable"]]
    if r["cable"].endswith("×2"):
        cables = [r["cable"][:-2]] * 2
    for c in cables:
        cab[c][p] += total
spare = Q["spare_len"] / 1000 * (1 + SLACK) * 6
cab["VVF2.0-3C"]["L-1"] += spare
cab_rows.append(("L1-13〜18", "2階予備（6回路）", "VVF2.0-3C×6", Q["spare_len"] / 1000 * 6, 0,
                 spare))
for c, v in cab.items():
    for p, m in v.items():
        count[("ケーブル", c, "600V ビニル絶縁ビニルシースケーブル平形")][p] += math.ceil(m)
for f in Q["feeders"]:
    p = "L-1" if "L-1" in f["name"] else "L-2"
    m = (f["plan"] + f["drop"]) / 1000 * (1 + SLACK)
    count[("幹線", f["cable"].split()[0], f["cable"].split()[1] + "（図示範囲。引込点からは別途加算）")][p] += math.ceil(m)
    cab_rows.append((f["name"], "幹線", f["cable"], f["plan"] / 1000, f["drop"] / 1000, m))

for key, pnl in PANELS.items():
    n_elb = sum(1 for c in pnl["circuits"] if c[2].startswith("ELB"))
    n_mcb = len(pnl["circuits"]) - n_elb
    count[("分電盤", pnl["name"], f"主幹 {pnl['main']} / 分岐 ELB×{n_elb}・MCB×{n_mcb}（{len(pnl['circuits'])}回路）")][key] += 1

UNIT = {"器具・配線器具": "台", "ボックス・プレート": "個", "ケーブル": "m", "幹線": "m", "分電盤": "面"}
UNIT_OVR = {"コンセント": "個", "片切スイッチ": "個", "3路スイッチ": "個", "確認表示灯内蔵スイッチ": "個",
            "スイッチプレート": "枚", "コンセントプレート": "枚"}

# ---------------------------------------------------------------- Excel
wb = Workbook()
ws = wb.active
ws.title = "数量総括表"
title(ws, "数量総括表　電灯・コンセント設備（1階 飲食店・事務所 複合店舗）",
      "電灯・コンセント設備平面図（1/50）より拾い出し。ケーブルは平面長＋立上下り＋余長10%。2026-10-04")
hdr = ["No", "区分", "品名", "仕様", "単位", "店舗 L-1", "事務所 L-2", "合計", "備考"]
ws.append([])
ws.append(hdr)
r0 = ws.max_row
ORDER = ["分電盤", "器具・配線器具", "ボックス・プレート", "ケーブル", "幹線"]
NOTE = {"コンセント": "", "ジョイントボックス": "分岐点数（差込形コネクタは一式）"}
no = 0
for sec in ORDER:
    keys = sorted([k for k in count if k[0] == sec], key=lambda k: (k[1], k[2]))
    for k in keys:
        no += 1
        unit = UNIT_OVR.get(k[1], UNIT[sec])
        c = count[k]
        ws.append([no, sec, k[1], k[2], unit, c.get("L-1", 0) or None, c.get("L-2", 0) or None, None,
                   NOTE.get(k[1], "")])
        rr = ws.max_row
        ws.cell(rr, 8).value = f"=SUM(F{rr}:G{rr})"
        if "未定" in k[2] or "専用" in k[2]:
            ws.cell(rr, 9).value = "機器未定のため仮設定" if "厨房" in k[2] else ws.cell(rr, 9).value
ws.append([None, "雑材", "差込形コネクタ・ステープル・支持材 等", "", "式", 1, 1, None, ""])
ws.cell(ws.max_row, 8).value = f"=SUM(F{ws.max_row}:G{ws.max_row})"
style_table(ws, r0, ws.max_row, len(hdr))
widths(ws, [5, 14, 24, 52, 6, 10, 11, 9, 28])
ws.freeze_panes = ws.cell(r0 + 1, 1)

# ---- 室別器具内訳
ws2 = wb.create_sheet("室別内訳")
title(ws2, "室別 器具・配線器具 内訳")
items = sorted({k for c in room_count.values() for k in c})
ws2.append([])
ws2.append(["室名"] + items + ["計"])
r0 = ws2.max_row
for room, c in room_count.items():
    ws2.append([room] + [c.get(i) for i in items] + [None])
    rr = ws2.max_row
    ws2.cell(rr, len(items) + 2).value = f"=SUM(B{rr}:{get_column_letter(len(items) + 1)}{rr})"
ws2.append(["合計"] + [None] * (len(items) + 1))
rr = ws2.max_row
for j in range(2, len(items) + 3):
    col = get_column_letter(j)
    ws2.cell(rr, j).value = f"=SUM({col}{r0 + 1}:{col}{rr - 1})"
style_table(ws2, r0, rr, len(items) + 2)
widths(ws2, [24] + [12] * (len(items) + 1))
for c in ws2[r0]:
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws2.row_dimensions[r0].height = 45

# ---- 配線数量内訳
ws3 = wb.create_sheet("配線数量内訳")
title(ws3, "配線数量内訳（回路別）", "平面長：図上の配線長／ホームランは器具→分電盤の直交距離。余長10%。単位 m")
ws3.append([])
ws3.append(["回路", "区分", "ケーブル", "平面長 m", "立上下り m", "余長込み m"])
r0 = ws3.max_row
for circuit, kind, cable, plan, drop, tot in sorted(cab_rows, key=lambda x: (x[0][:2], len(x[0]), x[0])):
    ws3.append([circuit, kind, cable, round(plan, 2), round(drop, 2), None])
    rr = ws3.max_row
    ws3.cell(rr, 6).value = f"=ROUND((D{rr}+E{rr})*1.1,2)" if kind != "2階予備（6回路）" else round(tot, 2)
ws3.append(["合計", "", "", None, None, None])
rr = ws3.max_row
for j in (4, 5, 6):
    col = get_column_letter(j)
    ws3.cell(rr, j).value = f"=SUM({col}{r0 + 1}:{col}{rr - 1})"
style_table(ws3, r0, rr, 6)
widths(ws3, [12, 16, 30, 11, 11, 12])
ws3.freeze_panes = ws3.cell(r0 + 1, 1)

# ---- 分電盤表
ws4 = wb.create_sheet("分電盤回路表")
title(ws4, "分電盤 回路表")
row = 3
for key, pnl in PANELS.items():
    rows = panel_rows(key, DEV)
    ws4.cell(row, 1, f"{pnl['name']}　主幹 {pnl['main']}　幹線 {pnl['feeder']}　{pnl['system']}").font = \
        Font(name=FONT, size=11, bold=True)
    row += 1
    hdr = ["回路", "負荷名称", "遮断器", "電線", "相", "容量 VA", "備考"]
    for j, h in enumerate(hdr, 1):
        ws4.cell(row, j, h)
    r0 = row
    for r in rows:
        row += 1
        for j, v in enumerate([r["cid"], r["name"], r["brk"], r["cable"], r["phase"], r["va"] or None, r["note"]], 1):
            ws4.cell(row, j, v)
    row += 1
    ws4.cell(row, 2, "合計")
    ws4.cell(row, 6, f"=SUM(F{r0 + 1}:F{row - 1})")
    style_table(ws4, r0, row, len(hdr))
    rr_, tt_ = phase_sum(rows)
    row += 1
    ws4.cell(row, 2, f"R相 {rr_:,.0f} VA / T相 {tt_:,.0f} VA　需要率 {pnl['demand']} → 最大相電流 ≒ "
                     f"{max(rr_, tt_) * pnl['demand'] / 100:.0f} A").font = Font(name=FONT, size=10)
    row += 3
widths(ws4, [9, 34, 26, 16, 6, 10, 18])

# ---- 照度計算書
ws5 = wb.create_sheet("照度計算書")
title(ws5, "照度計算書（光束法）　N = E × A ／（F × U × M）",
      "天井高 2.4m。照明率は反射率 70/50/10 の一般値（器具決定後にメーカー照明率表で確認）")
ws5.append([])
hdr = ["室名", "間口 m", "奥行 m", "面積 ㎡", "基準面", "光源高さ H m", "室指数 k", "器具", "光束 F lm",
       "照明率 U", "保守率 M", "設計照度 E lx", "必要台数", "設置台数", "計算照度 lx"]
ws5.append(hdr)
r0 = ws5.max_row
for c in Q["calc"]:
    ws5.append([c["room"], round(c["X"], 2), round(c["Y"], 2), c["A"], c["work"], c["H"], None, c["fix"], c["lm"],
                c["U"], c["M"], c["E"], None, c["N"], None])
    rr = ws5.max_row
    ws5.cell(rr, 7).value = f"=ROUND(B{rr}*C{rr}/(F{rr}*(B{rr}+C{rr})),2)"
    ws5.cell(rr, 13).value = f"=ROUND(L{rr}*D{rr}/(I{rr}*J{rr}*K{rr}),2)"
    ws5.cell(rr, 15).value = f"=ROUND(N{rr}*I{rr}*J{rr}*K{rr}/D{rr},0)"
style_table(ws5, r0, ws5.max_row, len(hdr))
widths(ws5, [22, 8, 8, 9, 14, 10, 9, 40, 10, 9, 9, 11, 10, 10, 11])
ws5.row_dimensions[r0].height = 32
ws5.append([])
ws5.append(["※ 室指数は主たる矩形の間口・奥行で算出。面積は不整形部分（カウンター通路の下駄箱前・階段前、ボックス席の"
            "格子部除外）を含む実面積。"])

path = os.path.join(OUT, "数量表_電灯コンセント設備.xlsx")
wb.save(path)
print("saved", path)
for k in sorted(count, key=lambda k: (ORDER.index(k[0]), k[1], k[2])):
    print(k, dict(count[k]))
