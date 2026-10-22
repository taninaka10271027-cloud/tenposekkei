"""樽屋 電気・機械設備 拾い出し表 (Excel) 生成."""
import os
import sys
import math
from collections import Counter, defaultdict
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import data as D
import gen_wire as W
import gen_sched as SC
import takeoff as T

OUT = os.environ.get('SEKKEI_OUT', os.getcwd() + '/out').rstrip('/') + '/設備拾い出し表.xlsx'

F_IN = PatternFill('solid', fgColor='FFF2CC')     # 入力セル(黄)
F_SEC = PatternFill('solid', fgColor='D9E1F2')    # 見出し
F_TOT = PatternFill('solid', fgColor='BDD7EE')    # 合計(青系)
F_HEAD = PatternFill('solid', fgColor='1F4E78')
FONT_HEAD = Font(name='游ゴシック', bold=True, color='FFFFFF', size=10)
FONT = Font(name='游ゴシック', size=10)
FONT_B = Font(name='游ゴシック', size=10, bold=True)
FONT_LINK = Font(name='游ゴシック', size=10, color='008000')   # 他シート参照(緑字)
THIN = Side(style='thin', color='999999')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

wb = Workbook()
K = SC.calc()
rows, prow, deg = T.run()
pipes = T.link_and_pipes()
ducts = T.ducts()
boxes = T.switch_boxes()

DW = '明細_配線'
DD = '明細_ダクト配管'

# =====================================================================  明細_配線
wd = wb.active
wd.title = DW
hd = ['回路', '階', 'ケーブル', '区分', '平面長(m)', '立上下(m)', '計(m)', '区間・根拠']
for c, h in enumerate(hd, 1):
    x = wd.cell(1, c, h)
    x.fill, x.font, x.border = F_HEAD, FONT_HEAD, BOX


def kubun(cir):
    if cir.startswith('P-'):
        return '動力'
    no = int(cir.split('-')[-1])
    if no in (1, 2, 3, 4, 24):
        return '電灯'
    return 'コンセント'


r = 2
allrows = rows + prow
for x in allrows:
    vals = [x['cir'], x['floor'], x['cable'], kubun(x['cir']), round(x['plan'], 2), round(x['vert'], 2), None, x['note']]
    for c, v in enumerate(vals, 1):
        wd.cell(r, c, v).font = FONT
        wd.cell(r, c).border = BOX
    wd.cell(r, 7, '=E%d+F%d' % (r, r)).border = BOX
    r += 1
# 空調 内外連絡線
for p in pipes:
    vals = [p['ac'], p['floor'], 'VVF2.0-3C', '空調連絡', round(p['link_len'] - 1.0 - T.M.RISE[p['ac']], 2),
            round(T.M.RISE[p['ac']] + 1.0, 2), None, '室内機→室外機 冷媒ルート同経路 (立上下・余長含む)']
    for c, v in enumerate(vals, 1):
        wd.cell(r, c, v).font = FONT
        wd.cell(r, c).border = BOX
    wd.cell(r, 7, '=E%d+F%d' % (r, r)).border = BOX
    r += 1
WD_LAST = r - 1
for c, w in zip('ABCDEFGH', [9, 5, 14, 10, 10, 10, 10, 46]):
    wd.column_dimensions[c].width = w
wd.freeze_panes = 'A2'
wd.auto_filter.ref = 'A1:H%d' % WD_LAST


def sumifs_cable(cable, floor, kb):
    return "=ROUND(SUMIFS('%s'!$G$2:$G$%d,'%s'!$C$2:$C$%d,\"%s\",'%s'!$B$2:$B$%d,%d,'%s'!$D$2:$D$%d,\"%s\"),1)" % (
        DW, WD_LAST, DW, WD_LAST, cable, DW, WD_LAST, floor, DW, WD_LAST, kb)


# =====================================================================  明細_ダクト配管
wp = wb.create_sheet(DD)
hd = ['区分', '階', '系統', '規格', '長さ(m)', '曲り(箇所)', '外壁貫通', '根拠']
for c, h in enumerate(hd, 1):
    x = wp.cell(1, c, h)
    x.fill, x.font, x.border = F_HEAD, FONT_HEAD, BOX
r = 2
for d_ in ducts:
    vals = ['ダクト', d_['floor'], d_['label'], 'φ%d' % d_['dia'], round(d_['len'], 2), d_['elbows'], 1,
            'M-0%d ダクトルート長+接続0.3m' % d_['floor']]
    for c, v in enumerate(vals, 1):
        wp.cell(r, c, v).font = FONT
        wp.cell(r, c).border = BOX
    r += 1
for p in pipes:
    for kind, spec, L in (('冷媒', '%s/%s' % (p['liq'], p['gas']), p['ref_len']), ('ドレン', p['drain'], p['drain_len'])):
        vals = [kind, p['floor'], p['ac'], spec, round(L, 2), p['corners'] if kind == '冷媒' else '', 1,
                'M-0%d 配管ルート長+%s' % (p['floor'], '立上下(機種別想定)' if kind == '冷媒' else '屋外排水端末')]
        for c, v in enumerate(vals, 1):
            wp.cell(r, c, v).font = FONT
            wp.cell(r, c).border = BOX
        r += 1
WP_LAST = r - 1
for c, w in zip('ABCDEFGH', [8, 5, 16, 16, 10, 10, 9, 40]):
    wp.column_dimensions[c].width = w
wp.freeze_panes = 'A2'


def sumifs_dd(col, kind, spec, floor):
    return "=ROUND(SUMIFS('%s'!$%s$2:$%s$%d,'%s'!$A$2:$A$%d,\"%s\",'%s'!$D$2:$D$%d,\"%s\",'%s'!$B$2:$B$%d,%d),1)" % (
        DD, col, col, WP_LAST, DD, WP_LAST, kind, DD, WP_LAST, spec, DD, WP_LAST, floor)


# =====================================================================  集計シート共通
HEAD = ['No', '名称', '規格・仕様', '単位', '1F', '2F', '屋外・共通', '設計数量', 'ロス率', '所要数量', '単価', '金額',
        '拾い根拠', '備考']
WIDTH = [5, 26, 40, 6, 9, 9, 9, 10, 7, 10, 11, 13, 34, 30]


class Sum:
    def __init__(self, title, sub):
        self.ws = wb.create_sheet(title)
        ws = self.ws
        ws['A1'] = '樽屋 電気・機械設備工事　拾い出し表　' + sub
        ws['A1'].font = Font(name='游ゴシック', size=14, bold=True)
        ws['A2'] = '天井高 2,400 / 階高 1F 3,000・2F 3,000 / 数量は成果品図面(DXF)の計測値。単価(黄色)は未記入。'
        ws['A2'].font = FONT
        for c, (h, w) in enumerate(zip(HEAD, WIDTH), 1):
            x = ws.cell(4, c, h)
            x.fill, x.font, x.border = F_HEAD, FONT_HEAD, BOX
            x.alignment = Alignment(horizontal='center', vertical='center')
            ws.column_dimensions[get_column_letter(c)].width = w
        ws.freeze_panes = 'C5'
        self.r = 5
        self.no = 0
        self.first = 5

    def sec(self, t):
        ws = self.ws
        ws.cell(self.r, 2, '【' + t + '】').font = FONT_B
        for c in range(1, 15):
            ws.cell(self.r, c).fill = F_SEC
            ws.cell(self.r, c).border = BOX
        self.r += 1

    def item(self, name, spec, unit, f1=None, f2=None, fc=None, loss=None, basis='', note=''):
        ws, r = self.ws, self.r
        self.no += 1
        vals = [self.no, name, spec, unit, f1, f2, fc]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            cell.font = FONT_LINK if isinstance(v, str) and v.startswith('=') and '!' in v else FONT
            cell.border = BOX
        ws.cell(r, 8, '=SUM(E%d:G%d)' % (r, r))
        ws.cell(r, 9, loss)
        ws.cell(r, 9).fill = F_IN
        ws.cell(r, 9).number_format = '0%'
        ws.cell(r, 10, '=IF(I%d="",H%d,ROUNDUP(H%d*(1+I%d),0))' % (r, r, r, r))
        ws.cell(r, 11).fill = F_IN
        ws.cell(r, 12, '=IF(K%d="","",J%d*K%d)' % (r, r, r))
        ws.cell(r, 13, basis)
        ws.cell(r, 14, note)
        for c in range(1, 15):
            ws.cell(r, c).border = BOX
            if c >= 8:
                ws.cell(r, c).font = FONT
            if c in (2, 3, 13, 14):
                ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical='center')
        for c in (5, 6, 7, 8, 10):
            ws.cell(r, c).number_format = '#,##0.0;-#,##0.0;""' if unit == 'm' and c != 10 else '#,##0;-#,##0;""'
        ws.cell(r, 11).number_format = '#,##0'
        ws.cell(r, 12).number_format = '#,##0'
        self.r += 1

    def total(self):
        ws, r = self.ws, self.r
        ws.cell(r, 2, '合計').font = FONT_B
        ws.cell(r, 12, '=SUM(L%d:L%d)' % (self.first, r - 1))
        ws.cell(r, 12).number_format = '#,##0;-#,##0;""'
        for c in range(1, 15):
            ws.cell(r, c).fill = F_TOT
            ws.cell(r, c).border = BOX
        ws.cell(r + 2, 2, '※黄色セル(ロス率・単価)のみ入力。金額は所要数量×単価で自動計算。').font = FONT
        ws.cell(r + 3, 2, '※1F/2F欄の緑字は明細シートからの自動集計。').font = FONT


# ---------------------------------------------------------------------------- 集計値
cnt = defaultdict(Counter)
for fl, L in ((1, D.L1F), (2, D.L2F)):
    for d in L:
        cnt[fl][d['t']] += 1
outc = defaultdict(Counter)
for fl, L in ((1, D.OUT1F), (2, D.OUT2F)):
    for o in L:
        outc[fl][o['t']] += 1
swc = defaultdict(Counter)
for fl, L in ((1, D.SW1F), (2, D.SW2F)):
    for s in L:
        swc[fl][s['k']] += 1
boxc = defaultdict(Counter)
for b in boxes:
    boxc[b['floor']]['1' if b['n'] <= 3 else '2' if b['n'] <= 6 else '3'] += 1
nl1 = sum(1 for c in D.L1_CIRCUITS if not c[6])
ne1 = sum(1 for c in D.L1_CIRCUITS if c[6])

# =====================================================================  1 幹線動力
s = Sum('1_幹線動力', '［幹線動力設備］')
s.sec('分電盤・動力盤')
s.item('電灯分電盤 L-1', '屋内露出形 鋼板製 W600×H1000×D160 1φ3W100/200V 主幹MCCB3P100AF/%dAT(中性線欠相保護付) 分岐MCB2P1E20A×%d、ELB2P2E20A30mA×%d' % (
    K['mainL'], nl1, ne1), '面', 1, None, None, None, 'E-07 単線結線図 / E-08 姿図', '盤メーカー見積')
s.item('動力盤 P-1', '屋内露出形 鋼板製 W500×H700×D160 3φ3W200V 主幹MCCB3P100AF/%dAT 分岐ELB3P30A30mA×3、ELB3P20A30mA×3' % K['mainP'],
       '面', 1, None, None, None, 'E-07 単線結線図 / E-08 姿図', '盤メーカー見積')
s.item('電力量計 取付', '電灯用(電力会社支給品)', '個', None, None, 1, None, 'E-07 幹線系統図', '支給品取付')
s.item('電力量計 取付', '動力用(電力会社支給品)', '個', None, None, 1, None, 'E-07 幹線系統図', '支給品取付・E-01で追加')
s.item('計器箱', '屋外形 W400×H440(参考)', '個', None, None, 2, None, 'E-08 機器姿図', '電力会社仕様確認')
s.item('手元開閉器', '屋外形 ELB箱 3P 30A 30mA (AC-1 室外機)', '個', None, None, 1, None, 'E-04 手元開閉器')
s.item('手元開閉器', '屋外形 ELB箱 3P 20A 30mA (AC-2・AC-3 室外機)', '個', None, None, 2, None, 'E-04 手元開閉器')
s.item('手元開閉器', '屋内形 ELB箱 3P 30A 30mA (いけす)', '個', 1, None, None, None, 'E-04 手元開閉器')
s.item('いけす用電源', '3P 30A 250V 接地極付 (端子又はコンセント)', '個', 1, None, None, None, 'E-04 P-1-1', '受口形状はいけす業者と確認')
s.sec('幹線')
s.item('幹線ケーブル', 'CV 38sq-3C (Wh→L-1)', 'm', None, None, 15, 0.1, 'E-07 幹線系統図 L≒15m(想定)', '引込口〜盤の実長で修正')
s.item('幹線ケーブル', 'CV 14sq-3C (Wh→P-1)', 'm', None, None, 15, 0.1, 'E-07 幹線系統図 L≒15m(想定)', '同上')
s.item('幹線端末処理', 'CV 38sq-3C 両端', '箇所', None, None, 2, None, 'E-07')
s.item('幹線端末処理', 'CV 14sq-3C 両端', '箇所', None, None, 2, None, 'E-07')
s.sec('動力分岐配線')
s.item('CVケーブル', 'CV 5.5sq-3C (P-1-1 いけす / P-1-2 AC-1)', 'm', sumifs_cable('CV5.5sq-3C', 1, '動力'), None, None, 0.1,
       '明細_配線 (E-04 ルート長+立上下)', '盤0.5m・外壁下り・手元開閉器余長1.0m含む')
s.item('CVケーブル', 'CV 3.5sq-3C (P-1-3 AC-2 / P-1-4 AC-3)', 'm', sumifs_cable('CV3.5sq-3C', 1, '動力'), None, None, 0.1,
       '明細_配線 (E-04)', '同上')
s.item('接地線', 'IV 2.0 (動力分岐 同ルート)', 'm', sumifs_cable('IV2.0', 1, '動力'), None, None, 0.1, '明細_配線 (E-04)')
s.item('外壁貫通', '動力ケーブル 室外機用 (防水処理共)', '箇所', None, None, 3, None, 'E-04 P-1-2〜4')
s.item('合成樹脂可とう電線管', 'PF22 屋外露出部保護 (外壁下り)', 'm', None, None, round(3 * (2.4 - 0.6 + 0.3), 1), 0.1,
       '室外機3台×(天井高2.4-端子FL+0.6+貫通0.3)', '耐候性PF')
s.sec('接地')
s.item('接地極', 'D種 接地棒 φ14×1.5m (L-1・P-1 各1)', '箇所', None, None, 2, None, 'E-07 幹線系統図 ED', '位置図示なし→要調整')
s.item('接地線', 'IV 5.5sq (盤→接地極)', 'm', None, None, round(2 * (5 + 1.9), 1), 0.1, '想定 水平5m+立下り1.9m ×2', '図示なし・想定')
s.item('接地抵抗測定・絶縁抵抗測定', '試験成績書提出共', '式', None, None, 1, None, '')
s.total()

# =====================================================================  2 電灯コンセント
s = Sum('2_電灯コンセント', '［電灯コンセント設備］')
s.sec('電灯配線 (照明・換気扇・非常照明)')
for cab, sp, nt in (('VVF2.0-2C', 'VVF 2.0mm-2C (電源・送り)', '電源配線「2.0」表記区間・ホームラン・立上り'),
                   ('VVF1.6-2C', 'VVF 1.6mm-2C', 'JB→器具・器具送り・スイッチ線'),
                   ('VVF1.6-3C', 'VVF 1.6mm-3C', 'スイッチ線3本以上・3路 (4本以上は2C/3Cの組合せ)')):
    s.item('VVFケーブル', sp, 'm', sumifs_cable(cab, 1, '電灯'), sumifs_cable(cab, 2, '電灯'), None, 0.1,
           '明細_配線 (E-03/E-05 条数別 ルート長+立上下)', nt + ' / SW立下り1.1m(厨房0.9m)')
s.sec('コンセント配線')
for cab, sp in (('VVF2.0-2C', 'VVF 2.0mm-2C'), ('VVF2.0-3C(E)', 'VVF 2.0mm-3C (接地線入 白黒緑)')):
    s.item('VVFケーブル', sp, 'm', sumifs_cable(cab, 1, 'コンセント'), sumifs_cable(cab, 2, 'コンセント'), None, 0.1,
           '明細_配線 (E-04/E-06 ルート長+立上下)', 'コンセント立下り2.1m(厨房0.9m・AC用0.4m)')
s.sec('配線器具')
s.item('埋込スイッチ', '片切 15A 300V (連用)', '個', swc[1]['S'], swc[2]['S'], None, None, 'E-03/E-05 ●', 'FL+1300(厨房FL+1500)')
s.item('埋込スイッチ', '3路 15A 300V (連用)', '個', swc[1]['3'], swc[2]['3'], None, None, 'E-03/E-05 ●3')
s.item('埋込コンセント', '2口 15A 125V', '個', outc[1]['2'], outc[2]['2'], None, None, 'E-04/E-06 「2」', 'FL+300')
s.item('埋込コンセント', '接地極付 2口 15A 125V', '個', outc[1]['2E'], outc[2]['2E'], None, None, 'E-04/E-06 「2E」')
s.item('埋込コンセント', '接地極付接地端子付 2口 15A 125V', '個', outc[1]['2ET'], outc[2]['2ET'], None, None, 'E-04 「2ET」 厨房専用', 'FL+1500')
s.item('埋込コンセント', 'エアコン用 接地極付接地端子付 15A/20A 250V', '個', outc[1]['EET 250V'], outc[2]['EET 250V'], None, None,
       'E-04/E-06 「EET250V」', 'FL+2000')
s.item('埋込コンセント', '接地極付 20A 125V (EV専用)', '個', outc[1]['2E EV'], None, None, None, 'E-04 「2E EV」', 'EV仕様確認')
s.item('埋込コンセント', '接地極付 15A 125V (AC-7 専用)', '個', None, outc[2]['2E AC'], None, None, 'E-06 「2E AC」', 'FL+2000')
s.item('防雨形コンセント', '接地極付 2口 15A 125V (WPE) ボックス共', '個', outc[1]['WPE'], None, None, None, 'E-04 「WPE」')
s.sec('ボックス・プレート')
nout1 = sum(v for k, v in outc[1].items() if k != 'WPE')
nout2 = sum(outc[2].values())
s.item('スイッチボックス', '1個用 (コンセント・スイッチ3連以下)', '個', boxc[1]['1'] + nout1, boxc[2]['1'] + nout2, None, None,
       'スイッチ群(200mm以内同一ボックス)+コンセント', '換気扇用コントロールスイッチ含む')
s.item('スイッチボックス', '2個用 (スイッチ4〜6連)', '個', boxc[1]['2'], boxc[2]['2'], None, None, '同上')
s.item('プレート', '1個用 (連用)', '個', boxc[1]['1'] + nout1, boxc[2]['1'] + nout2, None, None, 'ボックス数と同数')
s.item('プレート', '2個用 (連用)', '個', boxc[1]['2'], boxc[2]['2'], None, None, 'ボックス数と同数')
s.item('VVF用ジョイントボックス', '天井内 樹脂製 (図示⊘)', '個', T.JOINTS[(1, 'JB')], T.JOINTS[(2, 'JB')], None, None,
       'E-03/E-05 ⊘ (スイッチ結線点)')
s.item('差込形コネクタ', '2〜8本用 (JB・器具ジョイント部)', '個', T.CONNECT[1], T.CONNECT[2], None, 0.1,
       'JB・器具ジョイント部の電線(ネット)数', '器具送り部 %d+%d箇所含む' % (T.JOINTS[(1, '器具')], T.JOINTS[(2, '器具')]))
s.sec('2F予備・その他')
Lsp = (T.manh(D.RISER1F, T.PANEL_L_TOP) + (T.CEIL - T.PANEL_TOP) + T.FH[1]) / 1000
s.item('合成樹脂可とう電線管', 'PF22 2F予備回路用 空配管(呼線入)', 'm', round(Lsp * 4 - 3.0 * 4, 1), round(3.0 * 4, 1), None, 0.1,
       '盤→立上り+階高3.0m ×4本 (E-03/E-05注記)', 'L-1-20〜23')
s.item('ジョイントボックス', '2F天井内 予備配管端末', '個', None, 1, None, None, 'E-05 注記')
s.item('器具結線', '照明器具・換気扇・非常照明・誘導灯', '箇所', sum(cnt[1].values()) + len(D.FAN1F), sum(cnt[2].values()) + len(D.FAN2F),
       None, None, '照明器具・換気扇の台数', '器具代は 3_照明器具・4_換気')
s.item('配線用 支持・天井内ころがし', '天井内隠ぺい (支持材・ステップル)', '式', None, None, 1, None, '')
s.total()

# =====================================================================  3 照明器具
s = Sum('3_照明器具', '［照明器具設備］')
s.sec('照明器具 (照明プラン KOIZUMI 2026.10.18)')
for k in 'ABCDEFGHI':
    f = D.FIXTURES[k]
    s.item('照明器具 ' + k, '%s / %s / %s / %s' % (f['model'], f['name'], f['lamp'], f['mount']), '台', cnt[1][k] or None,
           cnt[2][k] or None, None, None, 'E-01/E-02 器具記号%s' % k, '%dVA' % f['va'])
s.sec('ライティングダクト (8000×1・5000×1・3000×1)')
for mdl, nm, n in (('WE0222', 'ライティングダクト本体 2000mm 黒', 2), ('WE0223', 'ライティングダクト本体 3000mm 黒', 4),
                   ('WE0241', 'フィードインキャップ', 3), ('WE0242', 'エンドキャップ', 3), ('WE0243', 'ジョイナーS', 3)):
    s.item(nm, 'KOIZUMI ' + mdl, '個', n, None, None, None, '照明プラン器具表 / E-01 LD', '1Fのみ')
s.item('ライティングダクト 吊り金具', '別途 (照明プラン注記「吊り金具別途」)', '式', 1, None, None, None, '照明プラン注記', '天井高2,400のため直付可否確認')
s.sec('非常用照明・誘導灯 (提案・所轄協議要)')
s.item('非常用照明', '電池内蔵形 LED 低天井用 (提案)', '台', cnt[1]['EM'], cnt[2]['EM'], None, None, 'E-01/E-02 ●非', '建基法施行令126条の4')
s.item('誘導灯', '避難口誘導灯 C級 LED 電池内蔵 (提案)', '台', cnt[1]['EX'], cnt[2]['EX'], None, None, 'E-01/E-02 避C', '消防法施行令26条')
s.total()
s.ws.cell(s.r + 4, 2, '※照明プランの照度計算は天井高3,000/2,500で算出。天井高2,400ではペンダント(C)の吊り長さ調整・照度再確認が必要。').font = FONT

# =====================================================================  4 換気
s = Sum('4_換気', '［換気設備］')
s.sec('換気機器 (換気機器明細書 T263764 / 機種選定書)')
s.item('ダクト用換気扇', 'VD-18ZVX7-C 天井埋込形 定風量 接続φ150 (VD-1/VD-2/VD-3)', '台', 2, 1, None, None, 'M-01/M-02 VD', '24h換気')
s.item('天吊金具', 'P-08TK3', '個', 2, 1, None, None, '機器明細書')
s.item('パイプ用ファン', 'V-12PFL8 24時間換気機能付 φ150 (PF-1/PF-2)', '台', 2, None, None, None, 'M-01 PF')
s.item('パイプ用ファン', 'V-08PSD8 接続φ100 (TF-1〜3 トイレ)', '台', 2, 1, None, None, 'M-01/M-02 TF')
s.item('自然給気グリル', 'P-18GLF6 φ150 (SA)', '個', len(D.GRILLE1F), len(D.GRILLE2F), None, None, 'M-01/M-02 SA')
hc = Counter()
for fl, lst in ((1, T.M.DUCT1F), (2, T.M.DUCT2F)):
    for lab, pts, dia, kind, hood in lst:
        if hood:
            hc[(fl, dia, hood[2] in T.M.ROAD_SIDE)] += 1
        elif lab.startswith('PF'):          # 既存プロットの北面フード
            hc[(fl, dia, False)] += 1
s.item('深形フード 防火ダンパー付', 'P-18VSD4 ステンレス φ150 FD付 (道路面・延焼ライン)', '個', hc[(1, 150, True)], hc[(2, 150, True)], None, None,
       'M-01/M-02 南面フード', '明細書P-18VS4から変更 (FD付)')
s.item('深形フード', 'P-18VS4 ステンレス φ150', '個', hc[(1, 150, False)], hc[(2, 150, False)], None, None, 'M-01/M-02 西・北面フード',
       'φ150計12個(明細書と一致)')
s.item('深形フード', 'P-13VS4 ステンレス φ100 (トイレ)', '個', hc[(1, 100, False)], hc[(2, 100, False)], None, None, 'M-01/M-02 北・東面フード',
       '東面も道路の場合はFD付に変更')
s.item('コントロールスイッチ', 'P-04SWLB5', '個', None, None, 2, None, '機器明細書', '階別割付は要確認')
s.item('コントロールスイッチ', 'P-04SWLV2', '個', None, None, 3, None, '機器明細書', '階別割付は要確認')
s.item('コントロールスイッチ', 'P-11SW2 (トイレ)', '個', 2, 1, None, None, '機器明細書')
s.item('換気扇 (倉庫)', '機種未定 EF-1/EF-2 + フード・スイッチ', '組', 2, None, None, None, 'E-01 既存プロット', '換気計画書に記載なし→要否確認')
s.sec('ダクト・付属品')
for dia in (150, 100):
    s.item('スパイラルダクト', 'φ%d 亜鉛鉄板' % dia, 'm', sumifs_dd('E', 'ダクト', 'φ%d' % dia, 1), sumifs_dd('E', 'ダクト', 'φ%d' % dia, 2),
           None, 0.1, '明細_ダクト配管 (M-01/M-02 ルート長+接続0.3m)')
    s.item('ダクト曲り (ベンド)', 'φ%d 90°' % dia, '個', sumifs_dd('F', 'ダクト', 'φ%d' % dia, 1), sumifs_dd('F', 'ダクト', 'φ%d' % dia, 2),
           None, None, '明細_ダクト配管 曲り箇所')
    s.item('外壁貫通', 'φ%d スリーブ・シール共' % dia, '箇所', sumifs_dd('G', 'ダクト', 'φ%d' % dia, 1), sumifs_dd('G', 'ダクト', 'φ%d' % dia, 2),
           None, None, '明細_ダクト配管', '倉庫換気扇(EF)分含む' if dia == 100 else '')
s.item('ダクト保温', 'グラスウール25t 外壁貫通部より1m', 'm', sum(1 for d in ducts if d['floor'] == 1), sum(1 for d in ducts if d['floor'] == 2),
       None, 0.1, 'ダクト系統数×1m', '結露防止 参考値')
s.item('防火ダンパー (区画貫通)', '防火区画貫通部 要否確認', '個', None, None, None, None, 'M-01/M-02 特記2', '外壁はFD付フードで対応済。区画は所轄確認')
s.total()
s.ws.cell(s.r + 4, 2, '※換気扇の電源配線は 2_電灯コンセント(電灯配線)に含む。厨房フード・排気ファン(火気使用室)は本計画外。').font = FONT

# =====================================================================  5 空調
s = Sum('5_空調', '［空調設備］')
s.sec('空調機 (三菱電機 空調プラン 2026.10.16)')
acs = D.AC1F + D.AC2F
cap = {'AC-1': '冷12.5/暖14.0kW', 'AC-2': '冷5.0/暖5.6kW', 'AC-3': '冷5.0/暖5.6kW', 'AC-4': '冷5.6/暖6.7kW', 'AC-5': '冷5.6/暖6.7kW',
       'AC-6': '冷5.6/暖6.7kW', 'AC-7': '冷3.6/暖4.2kW'}
groups = defaultdict(lambda: [0, 0, []])
for a in acs:
    fl = 1 if a in D.AC1F else 2
    key = (a['model'], a['kind'], a['pw'], cap[a['id']])
    groups[key][fl - 1] += 1
    groups[key][2].append(a['id'])
for (mdl, kind, pw, cp), (n1, n2, ids) in groups.items():
    s.item('空調機 (室内機+室外機)', '%s %s %s %s' % (mdl, kind, pw, cp), '組', n1 or None, n2 or None, None, None,
           'M-01/M-02 ' + '・'.join(ids), 'リモコン付属')
s.sec('冷媒配管・ドレン配管')
for pair in ('φ9.52/φ15.88', 'φ6.35/φ12.7', 'φ6.35/φ9.52'):
    s.item('冷媒配管', '被覆銅管 液%s/ガス%s (断熱材付) 2本1組' % tuple(pair.split('/')), 'm', sumifs_dd('E', '冷媒', pair, 1),
           sumifs_dd('E', '冷媒', pair, 2), None, 0.1, '明細_ダクト配管 (ルート長+立上下)')
for dr in ('VP25', 'VP20'):
    s.item('ドレン配管', '硬質塩ビ管 %s 保温付' % dr, 'm', sumifs_dd('E', 'ドレン', dr, 1), sumifs_dd('E', 'ドレン', dr, 2), None, 0.1,
           '明細_ダクト配管', 'AC-6は南面外壁へ直接排水')
s.item('内外連絡線', 'VVF 2.0mm-3C (室内機↔室外機)', 'm', sumifs_cable('VVF2.0-3C', 1, '空調連絡'), sumifs_cable('VVF2.0-3C', 2, '空調連絡'),
       None, 0.1, '明細_配線 (冷媒同ルート+余長1m)', 'メーカー指定サイズ確認')
s.item('外壁貫通 (配管)', 'スリーブ・パテ・シール共', '箇所', len(D.AC1F), len(D.AC2F), None, None, '各機1箇所')
s.item('配管化粧カバー', '屋外露出部', 'm', round(5 * 2.5, 1), round(2 * 1.0, 1), None, 0.1, '1F地上置き5台×2.5m / 2F屋根上2台×1.0m', '想定')
s.sec('据付・付帯')
s.item('室外機 据付', '地上置き 置台(プラ基礎)・アンカー', '台', len(D.AC1F), None, None, None, 'M-01 1F北面')
s.item('室外機 据付', '屋根上 防振架台・転倒防止', '台', None, len(D.AC2F), None, None, 'M-02 1F屋根上')
s.item('室内機 据付', '天吊・天井カセット (吊ボルト共)', '台', 3, None, None, None, 'AC-1〜3', 'カセットは天井開口(建築)')
s.item('室内機 据付', '壁掛形 (据付板共)', '台', 2, 2, None, None, 'AC-4〜7')
s.item('リモコン線', 'ワイヤードリモコン配線', '式', 1, None, None, None, 'AC-1〜3', 'リモコン位置未定・要確認')
s.item('気密試験・真空引き・冷媒追加充填', '冷媒追加量は配管実長により', '式', None, None, 1, None, '')
s.item('試運転調整', '', '式', None, None, 1, None, '')
s.total()
s.ws.cell(s.r + 4, 2, '※空調機の電源は 1_幹線動力(AC-1〜3 動力)・2_電灯コンセント(AC-4〜7 専用コンセント回路)に含む。').font = FONT

# =====================================================================  明細_器具
wm = wb.create_sheet('明細_器具')
hd = ['階', '種別', '記号/型', '回路', '点滅区分', 'X', 'Y', '取付高さ']
for c, h in enumerate(hd, 1):
    x = wm.cell(1, c, h)
    x.fill, x.font, x.border = F_HEAD, FONT_HEAD, BOX
r = 2


def put(vals):
    global r
    for c, v in enumerate(vals, 1):
        wm.cell(r, c, v).font = FONT
        wm.cell(r, c).border = BOX
    r += 1


for fl, L in ((1, D.L1F), (2, D.L2F)):
    for d in L:
        hh = 'FL+2000' if d['t'] in 'FGHI' else ('扉上' if d['t'] == 'EX' else '天井')
        put([fl, '照明', d['t'], d['cir'], d.get('grp') or '', d['x'], d['y'], hh])
for fl, L in ((1, D.SW1F), (2, D.SW2F)):
    for sw in L:
        k = {'S': '片切', '3': '3路', 'V': '換気SW'}[sw['k']]
        put([fl, 'スイッチ', k, '', sw['g'], sw['x'], sw['y'], 'FL+1500' if fl == 1 and T.in_kitchen(sw['x'], sw['y']) else 'FL+1300'])
for fl, L in ((1, D.OUT1F), (2, D.OUT2F)):
    for o in L:
        hh = 'FL+2000' if (o.get('ac') or 'AC' in o['t']) else ('FL+1500' if fl == 1 and T.in_kitchen(o['x'], o['y']) else 'FL+300')
        put([fl, 'コンセント', o['t'], o['cir'], '', o['x'], o['y'], hh])
for fl, L in ((1, D.FAN1F), (2, D.FAN2F)):
    for f in L:
        put([fl, '換気扇', f['id'] + ' ' + f['model'], f['cir'], f['g'], f['x'], f['y'], '天井'])
for c, w in zip('ABCDEFGH', [5, 10, 24, 9, 9, 9, 9, 10]):
    wm.column_dimensions[c].width = w
wm.freeze_panes = 'A2'

# =====================================================================  表紙・条件
wc = wb.create_sheet('表紙・条件', 0)
lines = [
    ('樽屋 電気・機械設備工事　拾い出し表', 16),
    ('作成日 2026年10月22日　／　出典：成果品図面 E-01〜E-08・M-01・M-02 (DXF)', 10),
    ('', 10),
    ('■ 計算条件', 12),
    ('ケーブル：電灯・コンセント＝VVF、幹線・動力分岐＝CV(接地線IV)。換気ダクト径は換気プラン凡例(φ150スパイラル)・機器仕様(V-08PSD8/P-13＝φ100)による', 10),
    ('天井高 2,400mm ／ 階高 1F 3,000mm・2F 3,000mm', 10),
    ('取付高さ：コンセントFL+300、スイッチFL+1300、厨房内FL+1500、ブラケットFL+2000、AC用コンセントFL+2000、盤上端FL+1900', 10),
    ('配線長 ＝ 図面上の配線ルート長(大R曲げ含む) ＋ 立上下(天井高−取付高) ＋ 階間3.0m(2F回路) ＋ 盤への配線(盤上端→天井0.5m)', 10),
    ('盤への配線(ホームラン)は回路の最寄り器具から盤までの直角距離で算定', 10),
    ('ダクト・配管長 ＝ M図のルート長 ＋ 接続・立上下 (冷媒は機種別立上下 2.0〜2.8m、ドレンは屋外排水端末1.0〜1.5m)', 10),
    ('ロス率：ケーブル・配管・ダクト 10%(初期値、黄色セルで変更可)。所要数量＝設計数量×(1+ロス率) 切上げ', 10),
    ('', 10),
    ('■ シート構成', 12),
    ('1_幹線動力 ／ 2_電灯コンセント ／ 3_照明器具 ／ 4_換気 ／ 5_空調 ：設備別 拾い出し表 (単価・金額欄付)', 10),
    ('明細_配線 ：回路・区間別のケーブル長 (1F/2F欄はここからSUMIFSで自動集計)', 10),
    ('明細_ダクト配管 ：ダクト・冷媒・ドレンの系統別長さ', 10),
    ('明細_器具 ：器具・スイッチ・コンセント・換気扇の全数リスト(座標・回路)', 10),
    ('', 10),
    ('■ セルの色', 12),
    ('黄色＝入力セル(ロス率・単価) ／ 緑字＝他シート参照 ／ 青塗り＝合計', 10),
    ('', 10),
    ('■ 注意・要確認事項', 12),
    ('1. 数量は図面の作図データからの計測値です。施工図・現地状況で必ず検算してください。単価は根拠資料が無いため未記入です。', 10),
    ('2. 幹線長 L≒15m は想定値(引込位置未確定)。接地極位置・接地線長も図示なしの想定値です。', 10),
    ('3. いけす容量は4.0kVA仮定。30A一杯の場合は P-1主幹75A・幹線22sq に変更(E-07注記)。', 10),
    ('4. 非常用照明・誘導灯は法規対応の提案数量。所轄協議で確定してください。', 10),
    ('5. 倉庫換気扇EF-1/2は機種未定。厨房フード・排気ファンは計画外(厨房業者確認)。', 10),
    ('6. 照明プランの照度計算は天井高3,000/2,500で実施。天井高2,400ではペンダント吊り長さ・照度の再確認が必要。', 10),
    ('7. ケーブル：電灯・コンセント＝VVF、幹線・動力＝CV (接地線 IV)。', 10),
    ('8. 空調機・換気機器の価格は各メーカー見積書(空調プラン・機器明細書)を参照。', 10),
    ('9. 道路面(延焼ライン)は南面と想定し、南面の外壁フード9個をFD付(P-18VSD4)とした。東面も道路の場合は要変更。', 10),
]
for i, (t, sz) in enumerate(lines, 1):
    c = wc.cell(i, 1, t)
    c.font = Font(name='游ゴシック', size=sz, bold=sz >= 12)
wc.column_dimensions['A'].width = 130
wc.cell(1, 1).fill = F_SEC

# 印刷設定
for ws in wb.worksheets:
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True

order = ['表紙・条件', '1_幹線動力', '2_電灯コンセント', '3_照明器具', '4_換気', '5_空調', DW, DD, '明細_器具']
wb._sheets = [wb[n] for n in order]
wb.active = 0
wb.save(OUT)
print('saved', OUT)
